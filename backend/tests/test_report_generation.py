import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api.v1.report_generator import get_report_generation_service
from app.core.exceptions import PropertySystemException, ResourceNotFoundError
from app.dependencies.auth import get_current_user
from app.main import app
from app.services import report_formats
from app.services.report_generation_service import (
    OPINION_PLACEHOLDER,
    ReportGenerationService,
    bank_template_context,
    render_docx,
    render_pdf,
)

USER = {"id": "u1", "role": "Reviewer"}
CASE = {"id": "case-1", "created_by": "u1", "property_name": "Wagholi Plot", "location": "Wagholi, Pune"}


class _DocService:
    def authorize_case_access(self, case_id, current_user):
        return CASE


class _DocRepo:
    def __init__(self, docs):
        self.docs = docs

    def list_by_case(self, case_id):
        return self.docs


class _Generic:
    def __init__(self, extractions):
        self.extractions = extractions

    def select_one(self, table, filters, order_by=None):
        return {"validated_json_output": self.extractions.get(filters["document_id"], {})}

    def select(self, table, filters=None, order_by=None):
        return [{"page_number": 1, "original_text": "page text", "english_text": ""}]


class _Flags:
    def __init__(self, flags):
        self.flags = flags

    def list_by_case(self, case_id):
        return self.flags


class _Builder:
    def __init__(self, excerpts):
        self.excerpts = excerpts

    def get_report_with_excerpts(self, case_id, current_user):
        return {"excerpts": self.excerpts}


DOCS = [
    {"id": "d1", "file_name": "sale_deed.pdf", "status": "approved"},
    {"id": "d2", "file_name": "seven_twelve.pdf", "status": "llm_done"},  # not approved: must be ignored
    {"id": "d3", "file_name": "ferfar.pdf", "status": "approved"},
]
EXTRACTIONS = {
    "d1": {
        "document_type": "sale_deed", "survey_number": "12/3", "owner_name": "Suresh Kale",
        "previous_owner_name": "Ramesh Patil", "property_location": "Wagholi", "area": "2 acres",
        "transaction_date": "2019-04-02", "registration_number": "REG-77",
        "chain": [
            {"order": 1, "owner": "Ramesh Patil", "date": "1998-01-01", "type": "Inheritance", "survey": "12/3"},
            {"order": 2, "owner": "Suresh Kale", "date": "2019-04-02", "type": "Sale Deed", "survey": "12/3"},
        ],
    },
    "d2": {"survey_number": "99/9", "owner_name": "Should Not Appear"},
    "d3": {"document_type": "mutation_record", "survey_number": "12/3",
           "chain": [{"order": 2, "owner": "Suresh Kale", "date": "2019-04-02", "type": "Sale Deed", "survey": "12/3"}]},
}
EXCERPTS = [
    {"document_name": "sale_deed.pdf", "document_id": "d1", "page_number": 2, "note": "Check seal",
     "excerpt_text": "Sold to Suresh\n| Plot | Area |\n| --- | --- |\n| A | 10 |"},
]
FLAGS = [{"flag_type": "boundary_mismatch", "severity": "high", "description": "East boundary differs", "status": "raised"}]


def _service(docs=DOCS, flags=FLAGS, excerpts=EXCERPTS, autofill=None):
    service = ReportGenerationService(_DocService(), _DocRepo(docs), _Generic(EXTRACTIONS), _Flags(flags), _Builder(excerpts))
    service._autofill = autofill or (lambda manifest, documents: {"fields": {}, "loops": {}})
    return service


def _flat(draft):
    parts = []
    for section in draft.sections:
        parts.append(section.heading)
        for block in section.blocks:
            parts.append(block.text)
            parts.extend(block.columns)
            for row in block.rows:
                parts.extend(row)
    return "\n".join(parts)


# --- formats ---

def test_formats_list_the_title_search_then_all_nine_banks():
    formats = report_formats.list_formats()
    assert formats[0]["id"] == "title-search"
    banks = formats[1:]
    assert len(banks) == 9 and all(f["kind"] == "bank" for f in banks)
    assert [f["label"].lower() for f in banks] == sorted(f["label"].lower() for f in banks)


def test_unknown_bank_format_is_rejected():
    with pytest.raises(ResourceNotFoundError):
        _service().build_draft("case-1", "../../etc/passwd", USER)


# --- title search draft ---

def test_title_search_uses_only_approved_documents_and_dedupes_the_chain():
    draft = _service().build_draft("case-1", "title-search", USER)
    text = _flat(draft)

    assert "sale_deed.pdf" in text and "ferfar.pdf" in text
    assert "seven_twelve.pdf" not in text and "Should Not Appear" not in text
    chain = next(s for s in draft.sections if s.id == "chain").blocks[0]
    assert [row[1] for row in chain.rows] == ["Ramesh Patil", "Suresh Kale"]  # Suresh Kale appears in two docs, listed once


def test_title_search_carries_selected_extracts_tables_and_a_visible_opinion_placeholder():
    draft = _service().build_draft("case-1", "title-search", USER)
    extracts = next(s for s in draft.sections if s.id == "extracts")

    assert any(b.type == "table" and b.rows == [["A", "10"]] for b in extracts.blocks)
    assert any("Note: Check seal" in b.text for b in extracts.blocks)
    opinion = next(s for s in draft.sections if s.id == "opinion")
    assert opinion.blocks[0].text == OPINION_PLACEHOLDER


def test_open_risk_flags_produce_a_warning():
    assert any("risk flag" in w for w in _service().build_draft("case-1", "title-search", USER).warnings)
    assert not any("risk flag" in w for w in _service(flags=[]).build_draft("case-1", "title-search", USER).warnings)


def test_no_approved_documents_blocks_the_draft():
    with pytest.raises(PropertySystemException) as error:
        _service(docs=[DOCS[1]]).build_draft("case-1", "title-search", USER)
    assert error.value.status_code == 409


# --- bank drafts ---

def test_bank_draft_prefills_facts_but_never_a_must_confirm_field():
    draft = _service().build_draft("case-1", "hdfc-bank-title-report", USER)
    fields = {f.name: f for f in draft.fields}

    assert fields["mortgagor_name"].value == "Suresh Kale"
    assert "Wagholi" in fields["property_description"].value
    must_confirm = [f for f in draft.fields if f.review_level == "must_confirm"]
    assert must_confirm and all(f.value == "" for f in must_confirm)
    assert next(l for l in draft.loops if l.name == "land_parcels").rows[0]["gat_number"] == "12/3"


@pytest.mark.parametrize("format_id", [f["id"] for f in report_formats.list_formats() if f["kind"] == "bank"])
def test_every_bank_template_renders_with_a_prefilled_draft(format_id):
    draft = _service().build_draft("case-1", format_id, USER)
    for field in draft.fields:
        if field.name in {"advocate_name"}:
            field.value = "Asha Rao"

    document = Document(io.BytesIO(render_docx(draft)))
    text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            text += "\n" + "\n".join(cell.text for cell in row.cells)

    assert "{{" not in text and "{%" not in text  # no unrendered template tags
    assert len(text.strip()) > 200


def test_bank_template_context_flattens_loops_and_drops_empty_rows():
    draft = _service().build_draft("case-1", "hdfc-bank-title-report", USER)
    flows = next(l for l in draft.loops if l.name == "title_flows")
    flows.rows = [{"gat_label": "Gat 800", "steps": "A. sold\nB. mutated"}, {"gat_label": "", "steps": []}]

    context = bank_template_context(draft)

    assert context["title_flows"] == [{"gat_label": "Gat 800", "steps": ["A. sold", "B. mutated"]}]


# --- files ---

def test_title_search_word_and_pdf_files_are_real_documents():
    draft = _service().build_draft("case-1", "title-search", USER)
    draft.sections[0].blocks[0].rows[0][1] = "Edited by lawyer"

    docx_text = "\n".join(p.text for p in Document(io.BytesIO(render_docx(draft))).paragraphs)
    pdf = render_pdf(draft)

    assert "Chain of title" in docx_text and "Opinion" in docx_text
    assert pdf.startswith(b"%PDF") and len(pdf) > 1000


def test_bank_pdf_without_libreoffice_says_so(monkeypatch):
    from app.services import report_generation_service as module

    monkeypatch.setattr(module, "_find_pdf_converter", lambda: None)
    draft = _service().build_draft("case-1", "hdfc-bank-title-report", USER)

    with pytest.raises(PropertySystemException) as error:
        render_pdf(draft)
    assert error.value.status_code == 501 and "Word" in error.value.message


# --- HTTP ---

def test_export_endpoint_streams_the_file_with_safe_headers():
    app.dependency_overrides[get_current_user] = lambda: USER
    app.dependency_overrides[get_report_generation_service] = lambda: _service()
    try:
        client = TestClient(app)
        draft = _service().build_draft("case-1", "title-search", USER).model_dump()
        response = client.post("/api/v1/cases/case-1/report/export", json={"draft": draft, "file_type": "docx"})
        formats = client.get("/api/v1/cases/case-1/report/formats")
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.headers["content-disposition"].startswith('attachment; filename="Title-Search-Report')
    assert response.headers["cache-control"] == "private, no-store"
    assert response.content[:2] == b"PK"  # a .docx is a zip
    assert formats.status_code == 200 and len(formats.json()) == 10


# --- autofill from the documents ---

def _manifest_fields(format_id="hdfc-bank-title-report"):
    return report_formats.get_manifest(format_id)["fields"]


def test_autofill_only_offers_facts_never_judgement_or_profile_fields():
    from app.services.report_autofill import eligible_fields

    names = {f["name"] for f in eligible_fields(_manifest_fields())}
    assert "opinion_conclusion" not in names and "advocate_name" not in names
    assert "mortgagor_name" in names or "property_description" in names


def test_autofilled_values_are_marked_and_judgement_fields_stay_empty_even_if_the_model_answers_them():
    manifest = report_formats.get_manifest("hdfc-bank-title-report")
    eligible_name = next(f["name"] for f in manifest["fields"] if f["source"] == "extracted" and f["name"] != "property_description"
                         and f["name"] not in {"mortgagor_name", "present_owner_names", "record_owner_name", "land_owner_names", "transferor_names", "gat_numbers"})

    def fake(manifest_arg, documents):
        assert [d["name"] for d in documents] == ["sale_deed.pdf", "ferfar.pdf"]  # approved only
        return {"fields": {eligible_name: "From the 7/12", "opinion_conclusion": "Title is clear"}, "loops": {}}

    draft = _service(autofill=fake).build_draft("case-1", "hdfc-bank-title-report", USER)
    fields = {f.name: f for f in draft.fields}

    assert fields[eligible_name].value == "From the 7/12" and fields[eligible_name].suggested
    assert fields["opinion_conclusion"].value == ""
    assert any("factual fields were filled" in w for w in draft.warnings)


def test_a_failed_autofill_still_gives_a_usable_form_and_says_so():
    def broken(manifest, documents):
        raise RuntimeError("model down")

    draft = _service(autofill=broken).build_draft("case-1", "hdfc-bank-title-report", USER)

    assert any("could not be read automatically" in w for w in draft.warnings)
    assert {f.name: f for f in draft.fields}["mortgagor_name"].value == "Suresh Kale"  # rule-based fill still applies


def test_clean_answer_drops_unknown_fields_bad_choices_and_empty_values():
    from app.services.report_autofill import clean_answer

    fields = [
        {"name": "village", "type": "text"},
        {"name": "mortgage_type", "type": "choice", "options": ["Simple", "Equitable"]},
        {"name": "taluka", "type": "text"},
    ]
    loops = [{"name": "land_parcels", "item_fields": ["gat_number", "east"], "nested_lists": {}}]
    raw = {
        "fields": {"village": " Wagholi ", "mortgage_type": "Something else", "taluka": None, "invented": "x"},
        "loops": {"land_parcels": [{"gat_number": "12/3", "east": "Road"}, {"gat_number": ""}, "junk"]},
    }

    assert clean_answer(raw, fields, loops) == {
        "fields": {"village": "Wagholi"},
        "loops": {"land_parcels": [{"gat_number": "12/3", "east": "Road"}]},
    }


def test_autofill_prompt_carries_only_approved_text_and_the_rules():
    from app.services.report_autofill import autofill

    seen = {}

    def fake(prompt):
        seen["prompt"] = prompt
        return '{"fields": {}, "loops": {}}'

    manifest = report_formats.get_manifest("hdfc-bank-title-report")
    autofill(manifest, [{"name": "a.pdf", "fields": {"owner_name": "X"}, "text": "hello"}], generate=fake)

    assert "Fill a field ONLY if an approved document clearly states it" in seen["prompt"]
    assert "opinion_conclusion" not in seen["prompt"] and "hello" in seen["prompt"]


def test_clean_answer_drops_non_answers_like_nil():
    from app.services.report_autofill import clean_answer

    fields = [{"name": "docs", "type": "text"}, {"name": "village", "type": "text"}]
    assert clean_answer({"fields": {"docs": "Nil", "village": "N/A"}}, fields, []) == {"fields": {}, "loops": {}}


@pytest.mark.parametrize("survey", [None, ""])
def test_a_report_can_be_drafted_for_a_case_with_no_survey_number(survey, monkeypatch):
    case = {**CASE, "survey_number": survey}
    service = _service()
    service.doc_service.authorize_case_access = lambda case_id, user: case
    draft = service.build_draft("case-1", "title-search", USER)
    assert draft.sections and draft.title == "Title Search Report"
